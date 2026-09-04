"""경인.GA7센터 월마감 리뷰 — 데이터 계산부.

한 달 마감과 다음 달 초반을 같이 놓고 보는 리뷰용. 대시보드(실시간 진척)와
목적이 다르다. 대시보드는 "지금 어디쯤인가"를 보고, 이건 "무엇을 얼마나
잘했고 그게 어떻게 만들어졌나"를 본다.

네 갈래 데이터를 합친다.
  · 인별실적       — 실적(천원). 마감월 최종본 + 당월 일별 스냅샷
  · GA매니저일보   — 설계 건수(전체→승인→체결). 실적의 선행지표
  · 9월매칭 취합본 — 중겸 본부(미라클·탑·리드·마스터) + 매니저 재배치
  · targets.json   — 센터별 월 목표

주의: 마감월 최종본은 파일 안 asOf 가 다음 달 1일이다(8월 최종본 asOf=9/1).
월 구분은 asOf 가 아니라 실적 컬럼 헤더('8월'/'9월')로 판단한다.
"""
from __future__ import annotations

import datetime
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import reader  # noqa: E402
import schema  # noqa: E402

RAW = Path("/Users/hoons/Documents/RAW DATA")
TARGETS = HERE / "targets.json"
CENTER = "경인.GA7센터"
DEPT = "경인컨설팅영업단"

# 마감월 최종본은 파일명이 스냅샷 규칙을 안 따른다. 명시적으로 짚어준다.
CLOSE_FILE = Path("/Users/hoons/Downloads/2608월최종(0901)_인별실적 (1).xlsx")
DAILY_FILE = RAW / "(공지)_26.08 마켓리딩본부 GA매니저일보_0831.xlsb"
MATCH_FILE = RAW / "(취합본)26년9월매칭_수도권마케팅2팀_경인영업단 (2).xlsx"

HQ_ORDER = ["Miracle 본부", "T.O.P 본부", "Lead 본부", "Master 본부"]
HQ_SHORT = {"Miracle 본부": "미라클", "T.O.P 본부": "탑", "Lead 본부": "리드",
            "Master 본부": "마스터", "#N/A": "미분류"}


# ---------------------------------------------------------------------------
# 월 판별 — asOf 가 아니라 실적 컬럼 헤더를 믿는다
# ---------------------------------------------------------------------------
def perf_month(path: Path) -> str | None:
    """실적 컬럼(17) 헤더에서 '8월' → '2026-08' 을 뽑는다."""
    import openpyxl
    from pyxlsb import open_workbook

    if path.suffix.lower() == ".xlsb":
        with open_workbook(reader._opened(path)) as wb:
            with wb.get_sheet("설계사") as sh:
                for i, row in enumerate(sh.rows()):
                    if i == 7:
                        cells = {c.c: c.v for c in row}
                        label, year = cells.get(17), None
                        break
                    if i == 2:
                        v = [c.v for c in row][0]
                        year = v
                else:
                    return None
    else:
        wb = openpyxl.load_workbook(reader._opened(path), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [r for i, r in zip(range(9), ws.iter_rows(values_only=True))]
        label = rows[7][17]
        year = rows[2][0]
        wb.close()

    mo = re.match(r"(\d{1,2})월", str(label or ""))
    if not mo:
        return None
    y = year.year if isinstance(year, datetime.datetime) else 2026
    month = int(mo[1])
    # 마감본은 asOf 가 다음 달 1일이라 연도가 넘어갈 수 있다(12월 마감 → asOf 1/1)
    if isinstance(year, datetime.datetime) and year.month == 1 and month == 12:
        y -= 1
    return f"{y}-{month:02d}"


def snapshots_for(month: str):
    """그 달 실적이 담긴 스냅샷들을 asOf 순으로."""
    out = []
    for p in reader.discover(RAW):
        if perf_month(p) == month:
            out.append(p)
    return sorted(out, key=lambda p: reader.parse_snapshot(p)["asOf"])


# ---------------------------------------------------------------------------
# 시상 종목 (종별)
# ---------------------------------------------------------------------------
def award_defs(path: Path):
    import openpyxl
    wb = openpyxl.load_workbook(reader._opened(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    raw = [r for i, r in zip(range(9), ws.iter_rows(values_only=True))]
    wb.close()
    rows = [{i: v for i, v in enumerate(r)} for r in raw]
    return schema.parse_awards(rows)


def award_kind(key: str) -> str:
    """시상 종목을 성격별로 묶는다 — 리뷰에서 '종별'로 쓰는 축."""
    if "월드투어" in key:
        return "월드투어"
    if "연속가동" in key:
        return "연속가동"
    if "신담보" in key:
        return "신담보"
    if "주력" in key:
        return "주력"
    if "인추가" in key:
        return "인(추가)"
    if re.match(r"\d주차\(인\)", key):
        return "인보험"
    return "기타"


# ---------------------------------------------------------------------------
# 중겸 본부 / 매니저
# ---------------------------------------------------------------------------
def load_match():
    import openpyxl
    wb = openpyxl.load_workbook(reader._opened(MATCH_FILE), read_only=True, data_only=True)

    s2 = [r for r in list(wb["2.사용인별 매칭"].iter_rows(values_only=True))[2:] if r[1]]
    mgr = {str(r[7]).strip().removesuffix(".0"): str(r[9]).strip() for r in s2}

    s1 = [r for r in list(wb["Sheet1"].iter_rows(values_only=True))[1:] if r[3]]
    hq = {}
    for r in s1:
        code = str(r[3]).strip().removesuffix(".0")
        v = str(r[8]).strip() if r[8] else ""
        hq[code] = v if v and v != "#N/A" else "#N/A"
    wb.close()
    return mgr, hq


# ---------------------------------------------------------------------------
# 매니저일보 — 설계 퍼널
# ---------------------------------------------------------------------------
def load_design():
    """{센터: {매니저: {'전체':[...], '승인':[...], '체결':[...], 'avg':..}}}, 날짜목록"""
    from pyxlsb import open_workbook

    with open_workbook(reader._opened(DAILY_FILE)) as wb:
        with wb.get_sheet("인별일자별설계") as sh:
            rows = [[c.v for c in r] for r in sh.rows()]

    dates = []
    for v in rows[1]:
        if isinstance(v, (int, float)) and v > 20000000:
            d = str(int(v))
            dates.append(f"{d[:4]}-{d[4:6]}-{d[6:8]}")
    first_col = next(i for i, v in enumerate(rows[1])
                     if isinstance(v, (int, float)) and v > 20000000)

    KIND = {"전체": "전체", "승인(%)": "승인", "체결(%)": "체결"}
    out = defaultdict(lambda: defaultdict(dict))
    for r in rows[5:]:
        if len(r) < 8 or not r[3] or r[6] not in KIND:
            continue
        center, mgr = str(r[3]).strip(), str(r[5]).strip()
        vals = [r[first_col + i] if first_col + i < len(r) else None
                for i in range(len(dates))]
        out[center][mgr][KIND[r[6]]] = [float(v) if isinstance(v, (int, float)) else 0.0
                                        for v in vals]
        if r[6] == "전체":
            out[center][mgr]["일평균"] = float(r[7]) if isinstance(r[7], (int, float)) else 0.0

    # 일보는 발행일 아침에 뽑혀서 마지막 날 컬럼이 통째로 비어 있다(8/31 파일 → 8/31 = 0).
    # 그대로 두면 일별 차트에 0 막대가 서고 영업일 수가 하루 부풀려진다.
    while dates:
        last = len(dates) - 1
        if any(v.get(k, [])[last] for mm in out.values() for v in mm.values()
               for k in ("전체",) if len(v.get(k, [])) > last):
            break
        dates.pop()
        for mm in out.values():
            for v in mm.values():
                for k in ("전체", "승인", "체결"):
                    if len(v.get(k, [])) > last:
                        v[k] = v[k][:last]
    return out, dates


# ---------------------------------------------------------------------------
# 조직 집계
# ---------------------------------------------------------------------------
def by_center(people):
    agg = defaultdict(lambda: {"cred": 0.0, "roster": 0, "active": 0, "dept": ""})
    for p in people:
        c = agg[p["center"]]
        c["cred"] += p["cred"]
        c["roster"] += 1
        c["active"] += 1 if p["cred"] > 0 else 0
        c["dept"] = p["dept"]
    return agg


def business_days(a: str, b: str):
    cur, end, out = datetime.date.fromisoformat(a), datetime.date.fromisoformat(b), []
    while cur <= end:
        if cur.weekday() < 5:
            out.append(cur.isoformat())
        cur += datetime.timedelta(days=1)
    return out


def collect():
    targets = json.loads(TARGETS.read_text(encoding="utf-8"))["months"]["2026-08"]["centers"]
    tgt = {k: v["revenue"] for k, v in targets.items()}

    # 매니저 표기는 원본 파일 값이 아니라 9월 확정 매칭을 따른다.
    # (사용인별 배정 → 지사 단위 override → 원본 순으로 해석한다.)
    close = reader.apply_current_manager_mapping(reader.parse_snapshot(CLOSE_FILE))
    aug_files = snapshots_for("2026-08")
    sep_files = snapshots_for("2026-09")

    aug_series, sep_series = [], []
    for p in aug_files:
        s = reader.apply_current_manager_mapping(reader.parse_snapshot(p))
        g = [x for x in s["people"] if x["center"] == CENTER]
        aug_series.append({"date": s["asOf"], "cred": sum(x["cred"] for x in g),
                           "active": sum(1 for x in g if x["cred"] > 0),
                           "centers": by_center(s["people"])})
    for p in sep_files:
        s = reader.apply_current_manager_mapping(reader.parse_snapshot(p))
        g = [x for x in s["people"] if x["center"] == CENTER]
        sep_series.append({"date": s["asOf"], "cred": sum(x["cred"] for x in g),
                           "active": sum(1 for x in g if x["cred"] > 0),
                           "centers": by_center(s["people"])})

    sep_last = (reader.apply_current_manager_mapping(reader.parse_snapshot(sep_files[-1]))
                if sep_files else None)

    return {
        "targets": tgt,
        "close": close,
        "closeCenters": by_center(close["people"]),
        "augSeries": aug_series,
        "sepSeries": sep_series,
        "sepLast": sep_last,
        "sepCenters": by_center(sep_last["people"]) if sep_last else {},
        "augAwards": award_defs(CLOSE_FILE),
        "sepAwards": award_defs(sep_files[-1]) if sep_files else [],
    }


if __name__ == "__main__":
    d = collect()
    g = d["closeCenters"][CENTER]
    t = d["targets"][CENTER]
    print(f"8월 마감 GA7: 실적 {g['cred']/10:,.1f}만 / 목표 {t/10:,.0f}만 = {g['cred']/t*100:.1f}%")
    print(f"   재적 {g['roster']}명 · 가동 {g['active']}명 ({g['active']/g['roster']*100:.1f}%)")
    rank = sorted(((v["cred"] / d["targets"][k] * 100, k) for k, v in d["closeCenters"].items()
                   if k in d["targets"]), reverse=True)
    pos = [i for i, (_, k) in enumerate(rank, 1) if k == CENTER][0]
    print(f"   달성률 전체 센터 {pos}/{len(rank)}위")
    print(f"8월 스냅샷 {len(d['augSeries'])}개, 9월 스냅샷 {len(d['sepSeries'])}개")
    print(f"9월 누적: {d['sepSeries'][-1]['cred']/10:,.1f}만 · 가동 {d['sepSeries'][-1]['active']}명")
    print(f"8월 시상종목 {len(d['augAwards'])}개 / 9월 {len(d['sepAwards'])}개")
